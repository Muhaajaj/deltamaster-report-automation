#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


def main():
    parser = argparse.ArgumentParser(
        description="Merge DeltaMaster TopM + Addison exports, calculate KPIs and export formatted Excel."
    )
    parser.add_argument("--topm", required=True, help="Path to TopM Excel export")
    parser.add_argument("--addison", required=True, help="Path to Addison Excel export")
    parser.add_argument("--out", default="Ergebnis_final_strukturiert.xlsx", help="Output Excel file path")
    args = parser.parse_args()

    topm_path = Path(args.topm)
    addison_path = Path(args.addison)
    out_path = Path(args.out)

    # ===============================================================
    # 1) READ TOPM EXCEL REPORT (last sheet, header=6)
    # ===============================================================
    df = pd.read_excel(topm_path, sheet_name=-1, header=6)

    # ===============================================================
    # 2) Rename first two columns to Hilfsmittel / Filiale and extract KSt
    # ===============================================================
    df = df.copy()
    df.rename(columns={df.columns[0]: "Hilfsmittel", df.columns[1]: "Filiale"}, inplace=True)
    df["KSt"] = df["Filiale"].astype(str).str[:5]

    # ===============================================================
    # 3) Filter Hilfsmittel
    # ===============================================================
    df_clean = df[~df["Hilfsmittel"].isin(["Alle Hilfsmittel", "08 - Einlagen"])].copy()

    # ===============================================================
    # 4) Row-level percent calculations
    # ===============================================================
    df_clean["(7) DB I % =\n(6) / (1)"] = (
        df_clean["(6) DB I =\n(1) - (5)"] / df_clean["(1) Umsatz-\nberechnung"]
    )
    df_clean["AP DB I % mit FP"] = df_clean["AP DB I mit FP"] / df_clean["(1) Umsatz-\nberechnung"]

    # ===============================================================
    # 5) Modifikationen logic (10 & 18)
    # ===============================================================
    df_clean["Modifikationen"] = np.where(
        df_clean["Hilfsmittel"].isin(["10 - Gehhilfen", "18 - Kranken-/ Behindertenfahrzeuge"]),
        df_clean["AP DB I mit FP"],
        df_clean["(6) DB I =\n(1) - (5)"],
    )
    df_clean["DB I % Modifikationen"] = df_clean["Modifikationen"] / df_clean["(1) Umsatz-\nberechnung"]

    # ===============================================================
    # 5a) Modifikationen 09 & 32 logic
    # ===============================================================
    df_clean["Modifikationen 09 & 32"] = df_clean["Modifikationen"]

    mask_0932 = df_clean["Hilfsmittel"].isin(
        ["09 - Elektrostimulationsgeräte", "32 - Therapeutische Bewegungsgeräte"]
    )

    df_clean.loc[mask_0932, "Modifikationen 09 & 32"] = (
        df_clean.loc[mask_0932, "(1) Umsatz-\nberechnung"] * 0.80
    )

    df_clean["DB I % Modifikationen 09 & 32"] = (
        df_clean["Modifikationen 09 & 32"] / df_clean["(1) Umsatz-\nberechnung"]
    )

    # ===============================================================
    # 6) Aggregate to KSt (sum numeric columns except pct columns)
    # ===============================================================
    pct_cols = [
        "(7) DB I % =\n(6) / (1)",
        "AP DB I % mit FP",
        "DB I % Modifikationen",
        "DB I % Modifikationen 09 & 32",
    ]
    sum_cols = [c for c in df_clean.select_dtypes(include="number").columns if c not in pct_cols]
    df_filiale = df_clean.groupby("KSt", as_index=False)[sum_cols].sum()

    # ===============================================================
    # 7) Recompute pct after aggregation (correct % logic)
    # ===============================================================
    df_filiale["(7) DB I % =\n(6) / (1)"] = (
        df_filiale["(6) DB I =\n(1) - (5)"] / df_filiale["(1) Umsatz-\nberechnung"]
    )
    df_filiale["AP DB I % mit FP"] = df_filiale["AP DB I mit FP"] / df_filiale["(1) Umsatz-\nberechnung"]
    df_filiale["DB I % Modifikationen"] = df_filiale["Modifikationen"] / df_filiale["(1) Umsatz-\nberechnung"]
    df_filiale["DB I % Modifikationen 09 & 32"] = (
        df_filiale["Modifikationen 09 & 32"] / df_filiale["(1) Umsatz-\nberechnung"]
    )

    # ===============================================================
    # 8) Representative Filiale per KSt (mode)
    # ===============================================================
    filiale_map = df_clean.groupby("KSt")["Filiale"].agg(lambda x: x.mode().iloc[0])
    df_filiale["Filiale"] = df_filiale["KSt"].map(filiale_map)

    # ===============================================================
    # 9) READ ADDISON EXCEL REPORT (last sheet, skiprows=8, header=None)
    # ===============================================================
    df2 = pd.read_excel(addison_path, sheet_name=-1, header=None, skiprows=8)

    # Map expected columns + extract KSt from Filiale
    df2 = df2.copy()
    df2.rename(columns={0: "Filiale", 2: "Art", 3: "Wert4", 5: "Wert6"}, inplace=True)
    df2["KSt"] = df2["Filiale"].astype(str).str[:5]

    # ===============================================================
    # 10) Filter relevant Arten + pivot
    # ===============================================================
    relevant_arten = ["Umsatzerlöse", "Aufwendungen für bez. Lfg. und Lst.", "Rohergebnis"]
    df_sub = df2[df2["Art"].isin(relevant_arten)]

    pivot4 = df_sub.pivot_table(index="KSt", columns="Art", values="Wert4", aggfunc="sum")
    pivot6 = df_sub.pivot_table(index="KSt", columns="Art", values="Wert6", aggfunc="sum")

    df2_new = pivot4.copy()
    # Use .get to avoid KeyError if a column is missing in synthetic/sample
    df2_new["Umsatzerlöse Kum"] = pivot6.get("Umsatzerlöse")
    df2_new["Aufwendungen für bez. Lfg. und Lst. Kum"] = pivot6.get("Aufwendungen für bez. Lfg. und Lst.")
    df2_new.reset_index(inplace=True)

    # ===============================================================
    # 11) Merge
    # ===============================================================
    df_merged = df_filiale.merge(df2_new, on="KSt", how="left")

    # ===============================================================
    # 12) Aufwendungen final
    # ===============================================================
    df_merged["Umsatzerlöse"] = df_merged.get("Umsatzerlöse", 0).fillna(0)
    df_merged["Aufwendungen für bez. Lfg. und Lst."] = df_merged.get("Aufwendungen für bez. Lfg. und Lst.", 0).fillna(0)

    df_merged["Aufwendungen final"] = (
        df_merged["Umsatzerlöse"] * (1 - df_merged["DB I % Modifikationen"]) +
        df_merged["Aufwendungen für bez. Lfg. und Lst."]
    ).round(0)

    # ===============================================================
    # 13) Format pct columns as strings
    # ===============================================================
    for col in [
        "(7) DB I % =\n(6) / (1)",
        "AP DB I % mit FP",
        "DB I % Modifikationen",
        "DB I % Modifikationen 09 & 32",
    ]:
        if col in df_merged.columns:
            df_merged[col] = df_merged[col].apply(lambda x: f"{x * 100:.1f}%" if pd.notnull(x) else "")

    # ===============================================================
    # 14) Final export column order
    # ===============================================================
    final_cols = [
        "KSt", "Filiale",
        "Aufträge",
        "(1) Umsatz-\nberechnung",
        "(2) Netto EK",
        "(3) Netto EK\nOhne WK",
        "(4) WK EK",
        "AP_EK_Verrechnung_WK_mit_FP",
        "(5) =\n(3) + (4)",
        "(6) DB I =\n(1) - (5)",
        "AP DB I mit FP",
        "Modifikationen",
        "Modifikationen 09 & 32",
        "(7) DB I % =\n(6) / (1)",
        "AP DB I % mit FP",
        "DB I % Modifikationen",
        "DB I % Modifikationen 09 & 32",
        "Umsatzerlöse",
        "Aufwendungen für bez. Lfg. und Lst.",
        "Rohergebnis",
        "Umsatzerlöse Kum",
        "Aufwendungen für bez. Lfg. und Lst. Kum",
        "Aufwendungen final",
    ]
    final_cols_existing = [c for c in final_cols if c in df_merged.columns]
    df_export = df_merged[final_cols_existing].copy()

    # ===============================================================
    # 15) Export Excel
    # ===============================================================
    df_export.to_excel(out_path, index=False, sheet_name="Auswertung")

    # ===============================================================
    # 16) Highlight columns in Excel
    # ===============================================================
    wb = load_workbook(out_path)
    ws = wb["Auswertung"]

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header = [cell.value for cell in ws[1]]
    columns_to_color = ["Aufwendungen final", "DB I % Modifikationen"]

    for col_name in columns_to_color:
        if col_name in header:
            col_idx = header.index(col_name) + 1
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                for cell in row:
                    cell.fill = yellow_fill

    wb.save(out_path)

    print(f"Done. Output saved to: {out_path.resolve()}")


if __name__ == "__main__":
    main()
